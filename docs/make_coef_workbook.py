"""P1 V2係数一覧Excelの生成スクリプト。  python docs/make_coef_workbook.py docs/P1_V2係数一覧.xlsx"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
H = PatternFill("solid", fgColor="1F3864")
HF = Font(bold=True, color="FFFFFF")
IN = PatternFill("solid", fgColor="FFF2CC")
GR = PatternFill("solid", fgColor="E7E6E6")
thin = Side(style="thin", color="BBBBBB")
B = Border(left=thin, right=thin, top=thin, bottom=thin)

# ---------------- 係数一覧 ----------------
ws = wb.active
ws.title = "係数一覧"
ws["A1"] = "V2 工数係数一覧（概算見積）  ※黄色の「変更値」に直したい値だけ入力。空欄は現行値を採用"
ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = ("工数 = 基本 + 台数 + 走行パターン + レイアウト(現地のみ) + 充電器 + 磁気停止 + ランドマーク + EV + 自動ドア + PLC信号 + 上位連携"
            " → 最低=基本工数、0.5h単位切上げ。金額 = 工数 × 賃率")
hdr = ["区分", "項目", "DBフィールド", "現行値", "変更値", "採用値", "単位", "意味・使われ方"]
for i, h in enumerate(hdr, 1):
    c = ws.cell(4, i, h)
    c.fill, c.font, c.border = H, HF, B
    c.alignment = Alignment(horizontal="center", vertical="center")

rows = [
    ("賃率", "原価賃率", "cost_hourly_rate", 6820, "円/h", "確定済み（工数×この値＝原価）"),
    ("賃率", "販売賃率", "sell_hourly_rate", 12000, "円/h", "確定済み（工数×この値＝販売額）"),
    ("基本", "設計 基本工数", "design_base_hours", 8, "h", "最低工数。どんな案件でもこれ以上"),
    ("基本", "現地 基本工数", "setup_base_hours", 8, "h", "最低工数"),
    ("台数", "設計 台数あたり", "design_vehicle_hours", 2, "h/台", "台数×この値を加算"),
    ("台数", "現地 台数あたり", "setup_vehicle_hours", 4, "h/台", "台数×この値を加算"),
    ("パターン", "設計 パターン点あたり", "design_pattern_hours", 3, "h/点", "走行パターンスコア×この値"),
    ("パターン", "現地 パターン点あたり", "setup_pattern_hours", 4, "h/点", "同上（現地デバッグ・検証）"),
    ("レイアウト", "現地 エリア規模あたり", "setup_layout_scale_hours", 1.5, "h/点", "√(短辺×長辺)÷10 × この値（現地のみ）"),
    ("ルート複雑さ", "simple", "route_factor_simple", 1.0, "倍", "単純な搬送"),
    ("ルート複雑さ", "standard", "route_factor_standard", 1.5, "倍", ""),
    ("ルート複雑さ", "complex", "route_factor_complex", 2.3, "倍", "分岐・交差が多い"),
    ("ルート複雑さ", "unknown", "route_factor_unknown", 1.6, "倍", "未確定時の安全側"),
    ("すれ違い・合流", "low", "merge_factor_low", 0.8, "倍", ""),
    ("すれ違い・合流", "standard", "merge_factor_standard", 1.3, "倍", ""),
    ("すれ違い・合流", "high", "merge_factor_high", 2.0, "倍", "デッドロック検証が増える"),
    ("すれ違い・合流", "unknown", "merge_factor_unknown", 1.4, "倍", ""),
    ("台数補正", "1台", "vehicle_factor_one", 1.0, "倍", "パターンスコアに乗算"),
    ("台数補正", "2台", "vehicle_factor_two", 1.4, "倍", ""),
    ("台数補正", "3台", "vehicle_factor_three", 1.8, "倍", ""),
    ("台数補正", "4台以上", "vehicle_factor_four_or_more", 2.2, "倍", "5台以上も同値（頭打ち）"),
    ("ランドマーク係数", "small", "landmark_factor_small", 1.0, "倍", "ランドマークを使う範囲"),
    ("ランドマーク係数", "medium", "landmark_factor_medium", 1.8, "倍", ""),
    ("ランドマーク係数", "large", "landmark_factor_large", 2.8, "倍", ""),
    ("ランドマーク係数", "full", "landmark_factor_full", 4.0, "倍", "全域"),
    ("ランドマーク係数", "unknown", "landmark_factor_unknown", 2.2, "倍", ""),
    ("ランドマーク工数", "設計 係数あたり", "design_landmark_hours", 1.0, "h/点", "ランドマーク係数×この値"),
    ("ランドマーク工数", "現地 係数あたり", "setup_landmark_hours", 4.0, "h/点", ""),
    ("設備", "設計 充電器1台", "design_charger_hours", 1.0, "h/台", ""),
    ("設備", "現地 充電器1台", "setup_charger_hours", 2.0, "h/台", ""),
    ("設備", "設計 磁気停止1箇所", "design_magnetic_stop_hours", 0.5, "h/箇所", ""),
    ("設備", "現地 磁気停止1箇所", "setup_magnetic_stop_hours", 1.0, "h/箇所", ""),
    ("設備連携", "設計 EV 1基", "design_elevator_hours", 4.0, "h/基", ""),
    ("設備連携", "現地 EV 1基", "setup_elevator_hours", 5.0, "h/基", ""),
    ("設備連携", "設計 自動ドア1基", "design_auto_door_hours", 2.0, "h/基", ""),
    ("設備連携", "現地 自動ドア1基", "setup_auto_door_hours", 3.0, "h/基", ""),
    ("設備連携", "設計 PLC信号1点", "design_plc_signal_hours", 1.2, "h/点", ""),
    ("設備連携", "現地 PLC信号1点", "setup_plc_signal_hours", 1.5, "h/点", ""),
    ("PLC規模→信号点数", "none", "plc_none_value", 0, "点", ""),
    ("PLC規模→信号点数", "small", "plc_small_value", 3, "点", "目安 1〜5点"),
    ("PLC規模→信号点数", "medium", "plc_medium_value", 10, "点", "目安 6〜15点"),
    ("PLC規模→信号点数", "large", "plc_large_value", 20, "点", "目安 16〜30点"),
    ("PLC規模→信号点数", "unknown", "plc_unknown_value", 6, "点", ""),
    ("上位連携(設計)", "none", "design_host_system_none_hours", 0, "h", ""),
    ("上位連携(設計)", "small", "design_host_system_small_hours", 4, "h", ""),
    ("上位連携(設計)", "medium", "design_host_system_medium_hours", 8, "h", ""),
    ("上位連携(設計)", "large", "design_host_system_large_hours", 16, "h", ""),
    ("上位連携(現地)", "none", "setup_host_system_none_hours", 0, "h", ""),
    ("上位連携(現地)", "small", "setup_host_system_small_hours", 4, "h", ""),
    ("上位連携(現地)", "medium", "setup_host_system_medium_hours", 8, "h", ""),
    ("上位連携(現地)", "large", "setup_host_system_large_hours", 16, "h", ""),
]
F = {}
r = 5
for g, n, f, v, u, m in rows:
    vals = [g, n, f, v, None, f'=IF(E{r}="",D{r},E{r})', u, m]
    for i, x in enumerate(vals, 1):
        ws.cell(r, i, x).border = B
    ws.cell(r, 5).fill = IN
    ws.cell(r, 4).fill = GR
    ws.cell(r, 6).font = Font(bold=True)
    F[f] = f"'係数一覧'!$F${r}"
    r += 1
for col, w in zip("ABCDEFGH", [18, 22, 32, 10, 10, 10, 8, 44]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"

# ---------------- 試算 ----------------
sc = wb.create_sheet("試算")
sc["A1"] = "工数・金額の試算（係数一覧の採用値で自動計算。黄色の入力は自由に変更可）"
sc["A1"].font = Font(bold=True, size=13)
sc["A2"] = "D〜F列はサンプル、G列は実案件用。シミュレーターから取れる値は右端「出力元」を参照"
heads = ["入力項目", "キー", "選択肢/単位", "小規模(例)", "中規模(例)", "大規模(例)", "実案件用", "シミュレーターの出力元"]
for i, h in enumerate(heads, 1):
    c = sc.cell(4, i, h)
    c.fill, c.font, c.border = H, HF, B
    c.alignment = Alignment(horizontal="center")

inp = [
    ("AGV台数", "vehicle_count", "台", [1, 3, 6, 2], "必要台数逆算（capacity_planner）"),
    ("受取ST数", "pickup_count", "箇所", [1, 2, 3, 1], "配置した受取ノード数"),
    ("受渡ST数", "dropoff_count", "箇所", [1, 2, 4, 1], "配置した受渡ノード数"),
    ("走行エリア 短辺", "area_short_m", "m", [10, 30, 60, 10], "走行エリア外形の短辺"),
    ("走行エリア 長辺", "area_long_m", "m", [20, 50, 100, 20], "走行エリア外形の長辺"),
    ("ルート複雑さ", "route_complexity", "simple/standard/complex/unknown", ["simple", "standard", "complex", "standard"], "分岐数・交差点数から判定（要ルール化）"),
    ("すれ違い・合流", "merge_level", "low/standard/high/unknown", ["low", "standard", "high", "standard"], "合流点・すれ違い点の数から判定（要ルール化）"),
    ("充電器数", "charger_count", "台", [0, 1, 2, 0], "配置した充電ノード数"),
    ("磁気停止数", "magnetic_stop_count", "箇所", [0, 0, 2, 0], "（手入力）"),
    ("ランドマーク範囲", "landmark_range", "small/medium/large/full/unknown", ["small", "medium", "large", "unknown"], "（手入力）"),
    ("EV数", "elevator_count", "基", [0, 0, 1, 0], "（手入力）"),
    ("自動ドア数", "auto_door_count", "基", [0, 1, 2, 0], "（手入力）"),
    ("PLC連携規模", "plc_scale", "none/small/medium/large/unknown", ["none", "small", "medium", "none"], "（手入力）"),
    ("上位システム連携", "host_system", "none/small/medium/large", ["none", "none", "small", "none"], "（手入力）"),
]
rowof = {}
for i, (lab, key, unit, vals, src) in enumerate(inp):
    rr = 5 + i
    rowof[key] = rr
    sc.cell(rr, 1, lab)
    sc.cell(rr, 2, key)
    sc.cell(rr, 3, unit)
    sc.cell(rr, 8, src)
    for j, v in enumerate(vals):
        sc.cell(rr, 4 + j, v).fill = IN
    for cc in range(1, 9):
        sc.cell(rr, cc).border = B

base = 5 + len(inp) + 1
sc.cell(base, 1, "▼ 中間値").font = Font(bold=True)
der = ["パターン数(受取×受渡)", "ルート係数", "合流係数", "台数係数", "パターンスコア", "レイアウトスコア",
       "ランドマーク係数", "PLC信号点数", "上位連携 設計h", "上位連携 現地h"]
dr = {n: base + 1 + i for i, n in enumerate(der)}
res0 = base + len(der) + 2
sc.cell(res0, 1, "▼ 結果").font = Font(bold=True)
resn = ["設計工数 (h)", "現地工数 (h)", "合計工数 (h)", "設計 原価(円)", "設計 販売(円)", "現地 原価(円)",
        "現地 販売(円)", "合計 原価(円)", "合計 販売(円)", "粗利率"]
rs = {n: res0 + 1 + i for i, n in enumerate(resn)}
for n, row in list(dr.items()) + list(rs.items()):
    sc.cell(row, 1, n)
    for cc in range(1, 8):
        sc.cell(row, cc).border = B


def lst(*keys):
    return ",".join(F[k] for k in keys)


for col in "DEFG":
    def g(k):
        return f"{col}${rowof[k]}"

    def d(k):
        return f"{col}{dr[k]}"

    def x(k):
        return f"{col}{rs[k]}"

    sc[d("パターン数(受取×受渡)")] = f"={g('pickup_count')}*{g('dropoff_count')}"
    sc[d("ルート係数")] = ('=CHOOSE(MATCH(%s,{"simple","standard","complex","unknown"},0),%s)'
                        % (g("route_complexity"), lst("route_factor_simple", "route_factor_standard", "route_factor_complex", "route_factor_unknown")))
    sc[d("合流係数")] = ('=CHOOSE(MATCH(%s,{"low","standard","high","unknown"},0),%s)'
                      % (g("merge_level"), lst("merge_factor_low", "merge_factor_standard", "merge_factor_high", "merge_factor_unknown")))
    n = g("vehicle_count")
    sc[d("台数係数")] = (f"=IF({n}<=1,{F['vehicle_factor_one']},IF({n}=2,{F['vehicle_factor_two']},"
                      f"IF({n}=3,{F['vehicle_factor_three']},{F['vehicle_factor_four_or_more']})))")
    sc[d("パターンスコア")] = f"=SQRT({d('パターン数(受取×受渡)')}*{d('ルート係数')}*{d('合流係数')})*{d('台数係数')}"
    sc[d("レイアウトスコア")] = f"=SQRT({g('area_short_m')}*{g('area_long_m')})/10"
    sc[d("ランドマーク係数")] = ('=CHOOSE(MATCH(%s,{"small","medium","large","full","unknown"},0),%s)'
                            % (g("landmark_range"), lst("landmark_factor_small", "landmark_factor_medium", "landmark_factor_large", "landmark_factor_full", "landmark_factor_unknown")))
    sc[d("PLC信号点数")] = ('=CHOOSE(MATCH(%s,{"none","small","medium","large","unknown"},0),%s)'
                         % (g("plc_scale"), lst("plc_none_value", "plc_small_value", "plc_medium_value", "plc_large_value", "plc_unknown_value")))
    for lab, pre in (("設計", "design"), ("現地", "setup")):
        sc[d(f"上位連携 {lab}h")] = ('=CHOOSE(MATCH(%s,{"none","small","medium","large"},0),%s)'
                                  % (g("host_system"), lst(*[f"{pre}_host_system_{k}_hours" for k in ("none", "small", "medium", "large")])))
    des = (f"{F['design_base_hours']}+{g('vehicle_count')}*{F['design_vehicle_hours']}+{d('パターンスコア')}*{F['design_pattern_hours']}"
           f"+{g('charger_count')}*{F['design_charger_hours']}+{g('magnetic_stop_count')}*{F['design_magnetic_stop_hours']}"
           f"+{d('ランドマーク係数')}*{F['design_landmark_hours']}+{g('elevator_count')}*{F['design_elevator_hours']}"
           f"+{g('auto_door_count')}*{F['design_auto_door_hours']}+{d('PLC信号点数')}*{F['design_plc_signal_hours']}+{d('上位連携 設計h')}")
    stp = (f"{F['setup_base_hours']}+{g('vehicle_count')}*{F['setup_vehicle_hours']}+{d('パターンスコア')}*{F['setup_pattern_hours']}"
           f"+{d('レイアウトスコア')}*{F['setup_layout_scale_hours']}"
           f"+{g('charger_count')}*{F['setup_charger_hours']}+{g('magnetic_stop_count')}*{F['setup_magnetic_stop_hours']}"
           f"+{d('ランドマーク係数')}*{F['setup_landmark_hours']}+{g('elevator_count')}*{F['setup_elevator_hours']}"
           f"+{g('auto_door_count')}*{F['setup_auto_door_hours']}+{d('PLC信号点数')}*{F['setup_plc_signal_hours']}+{d('上位連携 現地h')}")
    sc[x("設計工数 (h)")] = f"=CEILING(MAX({des},{F['design_base_hours']}),0.5)"
    sc[x("現地工数 (h)")] = f"=CEILING(MAX({stp},{F['setup_base_hours']}),0.5)"
    sc[x("合計工数 (h)")] = f"={x('設計工数 (h)')}+{x('現地工数 (h)')}"
    cr, sr = F["cost_hourly_rate"], F["sell_hourly_rate"]
    sc[x("設計 原価(円)")] = f"=ROUND({x('設計工数 (h)')}*{cr},0)"
    sc[x("設計 販売(円)")] = f"=ROUND({x('設計工数 (h)')}*{sr},0)"
    sc[x("現地 原価(円)")] = f"=ROUND({x('現地工数 (h)')}*{cr},0)"
    sc[x("現地 販売(円)")] = f"=ROUND({x('現地工数 (h)')}*{sr},0)"
    sc[x("合計 原価(円)")] = f"={x('設計 原価(円)')}+{x('現地 原価(円)')}"
    sc[x("合計 販売(円)")] = f"={x('設計 販売(円)')}+{x('現地 販売(円)')}"
    sc[x("粗利率")] = f"=IF({x('合計 販売(円)')}=0,0,({x('合計 販売(円)')}-{x('合計 原価(円)')})/{x('合計 販売(円)')})"
    sc[x("粗利率")].number_format = "0.0%"
    for k in ("設計 原価(円)", "設計 販売(円)", "現地 原価(円)", "現地 販売(円)", "合計 原価(円)", "合計 販売(円)"):
        sc[x(k)].number_format = "#,##0"
    for row in dr.values():
        sc[f"{col}{row}"].number_format = "0.00"
for k in ("設計工数 (h)", "現地工数 (h)", "合計工数 (h)", "合計 販売(円)"):
    for cc in range(1, 8):
        sc.cell(rs[k], cc).font = Font(bold=True)


def dv(key, opts):
    v = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=False)
    sc.add_data_validation(v)
    v.add(f"D{rowof[key]}:G{rowof[key]}")


dv("route_complexity", ["simple", "standard", "complex", "unknown"])
dv("merge_level", ["low", "standard", "high", "unknown"])
dv("landmark_range", ["small", "medium", "large", "full", "unknown"])
dv("plc_scale", ["none", "small", "medium", "large", "unknown"])
dv("host_system", ["none", "small", "medium", "large"])
for col, w in zip("ABCDEFGH", [24, 20, 30, 12, 12, 12, 12, 44]):
    sc.column_dimensions[col].width = w
sc.freeze_panes = "D5"

# ---------------- 反映手順 ----------------
nt = wb.create_sheet("反映手順")
for i, t in enumerate([
    "1. 『係数一覧』の黄色「変更値」に直したい値を入力（変えない項目は空欄のまま）",
    "2. 『試算』でサンプル案件の工数・金額が妥当か確認（妥当でなければ1へ）",
    "3. 確定したらこのファイルを渡す → 見積アプリの管理画面 /admin/v2-parameters とDBへ反映",
    "※ 賃率(原価6,820/販売12,000)は確定済みで反映済み",
    "※ ルート複雑さ・合流の判定ルール(シミュレーターのどの数値でsimple/standard/complexとするか)は別途決める",
], 1):
    nt.cell(i, 1, t)
nt.column_dimensions["A"].width = 110

wb.save(sys.argv[1])
